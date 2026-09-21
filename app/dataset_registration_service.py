from __future__ import annotations

import csv
import io
import re
from contextlib import contextmanager
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from fastapi import UploadFile
from openpyxl import load_workbook
from psycopg2.extras import Json, RealDictCursor

from app.dataset_ontology_mapping_service import get_ontology_fields
from app.db import get_connection
from app.endpoints.dataset_details import DatasetRegistryInput, _create_dataset_registry_from_payload

# Backs steps 1 ("Source & details") and 3 ("Review & publish") of the
# registration wizard. Step 2 (ontology mapping) already lives in
# dataset_ontology_mapping_service.py and is only called into here, not
# duplicated -- see suggest_mappings below.
#
# Source-type scope for this pass (product decision 2026-09-16): file
# (CSV/XLSX only, no PDF), url (test + detect JSON fields), database
# (.sql dump upload + static parse -- no live-credential connection, no
# execution of the dump). Map data (shapefile/GeoTIFF) is out of scope --
# that belongs to map-module-backend.
#
# --- v2 addition (2026-09-21) -------------------------------------------
# The v2 wizard (scratch-download/registration/v2) changes the product
# model: CPHR never takes a copy of the data, only a REFERENCE to wherever
# it already lives (a URI, a read-only connection string, or a map service
# layer), plus a "Verify reachability" check that reads just enough to
# drive ontology-mapping suggestions. See save_file_reference_source,
# save_database_reference_source, and save_map_service_source below --
# these are additive; the v1 upload-based file/url/database functions
# above are untouched.
#
# Scope decision (2026-09-21): live database verification ships for
# PostgreSQL only -- psycopg2 is the only DB driver in requirements.txt.
# Other engines are recorded but "Verify reachability" reports them as
# unsupported rather than guessing at a connection library.

REPO_ROOT = Path(__file__).resolve().parent.parent
UPLOAD_DIR = REPO_ROOT / "registration-uploads"


class DatasetNotFoundError(Exception):
    pass


class SourceConfigNotFoundError(Exception):
    pass


class UnsupportedFileTypeError(Exception):
    pass


@contextmanager
def _cursor():
    conn = get_connection()
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                yield cur
    finally:
        conn.close()


def _dataset_row(cur, dataset_id: int) -> dict | None:
    cur.execute("SELECT * FROM dataset_master WHERE dataset_id = %s", (dataset_id,))
    return cur.fetchone()


def _stringify(value) -> str | None:
    return None if value is None else str(value)


# ---------------------------------------------------------------------------
# Draft lifecycle
# ---------------------------------------------------------------------------

def create_draft(payload: DatasetRegistryInput) -> dict:
    result = _create_dataset_registry_from_payload(payload, dataset_status="draft")
    if result.get("status") == "error":
        raise ValueError(result["error"])
    return result


_DRAFT_UPDATABLE_COLUMNS = {
    "title", "description", "citation", "doi", "language", "data_language",
    "license", "keywords", "dataset_type", "category_id",
    "node_name", "node_maintained_by",
}


def update_draft(dataset_id: int, fields: dict) -> dict:
    """Patch a subset of dataset_master columns -- powers "Save as draft"."""
    updates = {k: v for k, v in fields.items() if k in _DRAFT_UPDATABLE_COLUMNS}
    if not updates:
        raise ValueError("No updatable fields supplied")

    set_clause = ", ".join(f"{col} = %s" for col in updates)
    values = list(updates.values()) + [dataset_id]

    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        cur.execute(
            f"UPDATE dataset_master SET {set_clause}, updated_at = NOW() WHERE dataset_id = %s RETURNING *",
            values,
        )
        row = cur.fetchone()
    return dict(row)


# ---------------------------------------------------------------------------
# Field detection -- pure parsing helpers (no DB, no disk) so they're
# testable directly against fixture content.
# ---------------------------------------------------------------------------

def parse_csv_fields(text: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    first_row = next(reader, {}) or {}
    fieldnames = reader.fieldnames or list(first_row.keys())
    return [
        {"field_name": name, "sample_value": _stringify(first_row.get(name))}
        for name in fieldnames
        if name
    ]


def parse_xlsx_fields(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header = next(rows, None) or ()
    first_data_row = next(rows, None) or ()

    fields = []
    for idx, name in enumerate(header):
        if not name:
            continue
        value = first_data_row[idx] if idx < len(first_data_row) else None
        fields.append({"field_name": str(name), "sample_value": _stringify(value)})
    return fields


_CREATE_TABLE_RE = re.compile(
    r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?[\"'`]?(?P<name>[\w.]+)[\"'`]?\s*\((?P<body>.*?)\)\s*;",
    re.IGNORECASE | re.DOTALL,
)
_COLUMN_LINE_RE = re.compile(r"""^["'`]?(?P<name>[A-Za-z_][\w]*)["'`]?\s+(?P<type>[A-Za-z][\w\s(),]*)""")
_SQL_CONSTRAINT_LINE_STARTS = ("primary key", "foreign key", "unique", "constraint", "check", "key ")


def _split_top_level_commas(body: str) -> list[str]:
    """Split a CREATE TABLE body on commas, ignoring commas nested in parens
    (e.g. inside a NUMERIC(10,2) type or a multi-column UNIQUE(...))."""
    parts, depth, current = [], 0, []
    for ch in body:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            parts.append("".join(current))
            current = []
        else:
            current.append(ch)
    if current:
        parts.append("".join(current))
    return parts


def parse_sql_dump_tables(sql_text: str) -> list[dict]:
    """Statically parse CREATE TABLE statements for table + column names.

    Never executes the dump -- an uploaded .sql file is untrusted input, and
    the registration flow only needs its schema, not its data.
    """
    tables = []
    for match in _CREATE_TABLE_RE.finditer(sql_text):
        table_name = match.group("name").split(".")[-1]
        columns = []
        for raw_line in _split_top_level_commas(match.group("body")):
            line = raw_line.strip()
            if not line or line.lower().startswith(_SQL_CONSTRAINT_LINE_STARTS):
                continue
            col_match = _COLUMN_LINE_RE.match(line)
            if col_match:
                columns.append({
                    "field_name": col_match.group("name"),
                    "data_type": col_match.group("type").strip().rstrip(","),
                })
        if columns:
            tables.append({"table_name": table_name, "columns": columns})
    return tables


def _save_upload(dataset_id: int, filename: str, content: bytes) -> Path:
    target_dir = UPLOAD_DIR / str(dataset_id)
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^\w.\-]", "_", filename)
    path = target_dir / safe_name
    path.write_bytes(content)
    return path


def _upsert_source_config(cur, dataset_id: int, **fields) -> dict:
    columns = ["dataset_id"] + list(fields.keys())
    values = [dataset_id] + list(fields.values())
    placeholders = ", ".join(["%s"] * len(values))
    update_clause = ", ".join(f"{c} = EXCLUDED.{c}" for c in fields)
    cur.execute(
        f"""
        INSERT INTO dataset_source_config ({", ".join(columns)})
        VALUES ({placeholders})
        ON CONFLICT (dataset_id) DO UPDATE SET {update_clause}, updated_at = NOW()
        RETURNING *
        """,
        values,
    )
    return dict(cur.fetchone())


async def save_file_source(dataset_id: int, upload: UploadFile) -> dict:
    filename = upload.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    content = await upload.read()

    if ext == "csv":
        fields = parse_csv_fields(content.decode("utf-8", errors="replace"))
    elif ext in {"xlsx", "xlsm"}:
        fields = parse_xlsx_fields(content)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type '.{ext or '?'}' -- only .csv and .xlsx are supported")

    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        stored_path = _save_upload(dataset_id, filename, content)
        return _upsert_source_config(
            cur, dataset_id,
            source_type="file",
            file_name=filename,
            file_path=str(stored_path),
            file_format=ext,
            detected_fields=Json(fields),
        )


def _parse_auth_header(auth_header: str | None) -> dict:
    if not auth_header:
        return {}
    if ":" in auth_header:
        key, _, value = auth_header.partition(":")
        return {key.strip(): value.strip()}
    return {"Authorization": auth_header}


async def test_and_save_url_source(
    dataset_id: int, source_url: str, method: str, auth_header: str | None, response_format: str
) -> dict:
    import httpx

    async with httpx.AsyncClient(timeout=10.0) as client:
        response = await client.request(method.upper(), source_url, headers=_parse_auth_header(auth_header))
        response.raise_for_status()
        data = response.json()

    if isinstance(data, list):
        sample = data[0] if data else {}
    elif isinstance(data, dict):
        sample = data
    else:
        sample = {}

    fields = [
        {"field_name": key, "sample_value": _stringify(value)}
        for key, value in sample.items()
    ] if isinstance(sample, dict) else []

    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        return _upsert_source_config(
            cur, dataset_id,
            source_type="url",
            source_url=source_url,
            http_method=method.upper(),
            auth_header=auth_header,
            response_format=response_format,
            detected_fields=Json(fields),
        )


async def save_database_dump_source(dataset_id: int, upload: UploadFile) -> dict:
    filename = upload.filename or "dump.sql"
    content = await upload.read()
    tables = parse_sql_dump_tables(content.decode("utf-8", errors="replace"))
    if not tables:
        raise ValueError("No CREATE TABLE statements were found in the uploaded dump")

    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        stored_path = _save_upload(dataset_id, filename, content)
        source_config = _upsert_source_config(
            cur, dataset_id,
            source_type="database",
            dump_file_name=filename,
            dump_file_path=str(stored_path),
            selected_table_name=None,
            detected_fields=Json({"tables": tables}),
        )
    return {"tables": tables, "source_config": source_config}


def select_database_table(dataset_id: int, table_name: str) -> dict:
    with _cursor() as cur:
        cur.execute("SELECT * FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,))
        config = cur.fetchone()
        if not config:
            raise SourceConfigNotFoundError(dataset_id)

        tables = (config["detected_fields"] or {}).get("tables", [])
        match = next((t for t in tables if t["table_name"] == table_name), None)
        if not match:
            raise ValueError(f"Table {table_name!r} was not found in the uploaded dump")

        cur.execute(
            """
            UPDATE dataset_source_config
            SET selected_table_name = %s, detected_fields = %s, updated_at = NOW()
            WHERE dataset_id = %s
            RETURNING *
            """,
            (table_name, Json(match["columns"]), dataset_id),
        )
        return dict(cur.fetchone())


# ---------------------------------------------------------------------------
# Ontology mapping auto-suggest (step 2) -- reads detected_fields saved
# above, matches against the real ontology field list from
# dataset_ontology_mapping_service.get_ontology_fields.
# ---------------------------------------------------------------------------

def _normalize_field_name(name: str) -> str:
    return re.sub(r"[\s_\-]", "", name or "").lower()


def _best_match(field_name: str, ontology_fields: list[dict]) -> tuple[dict | None, float]:
    """Pure matching logic, split out from suggest_mappings so it's testable
    without a database: highest SequenceMatcher ratio against either the
    ontology field's own name or its display label."""
    norm_field = _normalize_field_name(field_name)
    best_field, best_score = None, 0.0
    for candidate in ontology_fields:
        score = max(
            SequenceMatcher(None, norm_field, _normalize_field_name(candidate["value"])).ratio(),
            SequenceMatcher(None, norm_field, _normalize_field_name(candidate["label"])).ratio(),
        )
        if score > best_score:
            best_field, best_score = candidate, score
    return best_field, best_score

# Below this ratio, don't suggest anything -- an empty dropdown beats a
# confidently wrong one. Between this and AUTO_THRESHOLD, still suggest it
# but flag it for the "needs review" badge instead of "auto".
SUGGEST_THRESHOLD = 0.6
AUTO_THRESHOLD = 0.85


def suggest_mappings(dataset_id: int, ontology_graph_key: str) -> dict:
    with _cursor() as cur:
        cur.execute(
            "SELECT detected_fields FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,)
        )
        config = cur.fetchone()
    if not config:
        raise SourceConfigNotFoundError(dataset_id)

    detected = config["detected_fields"]
    detected = detected if isinstance(detected, list) else []

    ontology_fields = get_ontology_fields(ontology_graph_key)["items"]

    items = []
    for entry in detected:
        field_name = entry.get("field_name")
        if not field_name:
            continue
        best_field, score = _best_match(field_name, ontology_fields)
        items.append({
            "field_name": field_name,
            "suggested_ontology_field": best_field["value"] if best_field and score >= SUGGEST_THRESHOLD else None,
            "confidence": round(score, 2),
            "auto": score >= AUTO_THRESHOLD,
        })

    return {"dataset_id": dataset_id, "ontology_graph_key": ontology_graph_key, "items": items}


# ---------------------------------------------------------------------------
# Propose a new ontology term (step 2 escape hatch)
# ---------------------------------------------------------------------------

def propose_term(
    dataset_id: int, field_name: str, proposed_label: str,
    proposed_definition: str | None, target_ontology_graph_key: str | None,
) -> dict:
    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        cur.execute(
            """
            INSERT INTO ontology_term_proposal
                (dataset_id, field_name, proposed_label, proposed_definition, target_ontology_graph_key)
            VALUES (%s, %s, %s, %s, %s)
            RETURNING *
            """,
            (dataset_id, field_name, proposed_label, proposed_definition, target_ontology_graph_key),
        )
        return dict(cur.fetchone())


def list_term_proposals(status: str | None = None) -> list[dict]:
    with _cursor() as cur:
        if status:
            cur.execute(
                "SELECT * FROM ontology_term_proposal WHERE status = %s ORDER BY created_at DESC", (status,)
            )
        else:
            cur.execute("SELECT * FROM ontology_term_proposal ORDER BY created_at DESC")
        return [dict(r) for r in cur.fetchall()]


# ---------------------------------------------------------------------------
# Step 3 -- review & publish
# ---------------------------------------------------------------------------

def _reference_type_label(source_type: str | None, reference_uri: str | None = None) -> str | None:
    """Human-readable "Reference type" line for the review/success screen,
    e.g. "File path (s3://)" -- matching the v2 mockup's summary list."""
    if source_type == "file":
        scheme = _reference_uri_scheme(reference_uri) if reference_uri else ""
        return f"File path ({scheme}://)" if scheme else "File path"
    return {
        "url": "URL / API",
        "database": "Database",
        "map_service": "Map service",
    }.get(source_type)


def _reference_location(source: dict | None) -> str | None:
    """The one value worth showing back to the contributor for "Data
    location" -- never the raw db_connection_string (that stays server-side
    even here; the table/view name is the useful, non-sensitive part)."""
    if not source:
        return None
    return {
        "file": source.get("reference_uri") or source.get("file_path"),
        "url": source.get("source_url"),
        "database": source.get("selected_table_name"),
        "map_service": source.get("map_layer_name"),
    }.get(source.get("source_type"))


def get_review_summary(dataset_id: int) -> dict:
    with _cursor() as cur:
        dataset = _dataset_row(cur, dataset_id)
        if not dataset:
            raise DatasetNotFoundError(dataset_id)

        cur.execute(
            "SELECT category_name FROM category_master WHERE category_id = %s", (dataset["category_id"],)
        )
        category_row = cur.fetchone()

        cur.execute("SELECT * FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,))
        source = cur.fetchone()

        cur.execute("SELECT COUNT(*) AS count FROM dataset_mapping WHERE dataset_id = %s", (dataset_id,))
        fields_mapped = cur.fetchone()["count"]

    detected = source["detected_fields"] if source else None
    fields_detected = len(detected) if isinstance(detected, list) else 0

    return {
        "dataset_id": dataset_id,
        "title": dataset["title"],
        "category_id": dataset["category_id"],
        "category_name": category_row["category_name"] if category_row else None,
        "status": dataset["status"],
        "node_name": dataset.get("node_name"),
        "node_maintained_by": dataset.get("node_maintained_by"),
        "source_type": source["source_type"] if source else None,
        "reference_type_label": _reference_type_label(
            source["source_type"] if source else None,
            source.get("reference_uri") if source else None,
        ),
        "reference_location": _reference_location(source),
        "reachability_status": source.get("reachability_status") if source else None,
        "fields_detected": fields_detected,
        "fields_mapped": fields_mapped,
    }


def publish(dataset_id: int) -> dict:
    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        cur.execute(
            "UPDATE dataset_master SET status = %s, updated_at = NOW() WHERE dataset_id = %s RETURNING *",
            ("Pending review", dataset_id),
        )
        row = dict(cur.fetchone())

    try:
        from app.retrieval_index import build_index
        build_index()
    except Exception:
        pass  # best-effort reindex -- publish must succeed even if this fails

    return row


# ---------------------------------------------------------------------------
# v2 -- reference-based sources ("Point to your data"). Each save_* function
# stores where the data lives; each verify_* function is the "Verify
# reachability" button -- it reads just enough to confirm the reference is
# real and, where practical, detect fields for step 2's mapping suggestions.
# Nothing here ever copies the contributor's underlying records.
# ---------------------------------------------------------------------------

def _reference_uri_scheme(uri: str) -> str:
    """Lowercase URI scheme, or '' for a bare filesystem path (no scheme)."""
    return urlsplit(uri or "").scheme.lower()


def _split_schema_table(table_name: str) -> tuple[str, str]:
    """'public.plot_observations' -> ('public', 'plot_observations').

    A bare table name with no schema defaults to 'public', matching
    Postgres's own default search_path behaviour.
    """
    parts = (table_name or "").strip().split(".", 1)
    return (parts[0], parts[1]) if len(parts) == 2 else ("public", parts[0])


def _redact_connection_string(conn_str: str | None) -> str | None:
    """Mask the password in a connection string before it's ever echoed back
    over the API -- the frontend only needs to confirm what it already
    submitted, not have the password reflected at it again."""
    if not conn_str:
        return conn_str
    parts = urlsplit(conn_str)
    if not parts.password:
        return conn_str
    netloc = parts.netloc.replace(f":{parts.password}@", ":****@", 1)
    return urlunsplit((parts.scheme, netloc, parts.path, parts.query, parts.fragment))


def _public_source_config(row: dict) -> dict:
    """Return a source_config row safe to hand to the frontend -- with any
    stored connection-string password redacted."""
    row = dict(row)
    if "db_connection_string" in row:
        row["db_connection_string"] = _redact_connection_string(row["db_connection_string"])
    return row


def _capabilities_request_params(layer_type: str | None) -> dict:
    """Map a wizard layer-type choice to the OGC GetCapabilities query it
    implies. Accepts either the mockup's display label ("Vector (WFS)") or
    a bare service code ("WFS") so the frontend doesn't have to translate."""
    normalized = (layer_type or "").upper()
    if "WFS" in normalized:
        service = "WFS"
    elif "WMTS" in normalized:
        service = "WMTS"
    else:
        service = "WMS"
    return {"service": service, "request": "GetCapabilities"}


def save_file_reference_source(
    dataset_id: int, reference_uri: str, file_format: str | None, access_credentials_ref: str | None
) -> dict:
    """Step 2 (file path): record where the file already lives. Never
    fetched in full here -- verify_file_reference reads a header only."""
    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        return _upsert_source_config(
            cur, dataset_id,
            source_type="file",
            reference_uri=reference_uri,
            file_format=file_format,
            access_credentials_ref=access_credentials_ref,
            reachability_status="unverified",
            reachability_checked_at=None,
            reachability_detail=None,
        )


async def verify_file_reference(dataset_id: int) -> dict:
    """"Verify reachability" for a file reference. http(s) URIs get a
    ranged GET to sniff a header row; other schemes (s3://, nfs://, smb://,
    or a bare local path) can't be resolved from this process without
    scheme-specific credentials/drivers we don't have, so they're reported
    as 'unverified' with an explanation rather than guessed at."""
    with _cursor() as cur:
        cur.execute("SELECT * FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,))
        config = cur.fetchone()
    if not config or not config.get("reference_uri"):
        raise SourceConfigNotFoundError(dataset_id)

    uri = config["reference_uri"]
    scheme = _reference_uri_scheme(uri)
    detected_fields = None

    if scheme in ("http", "https"):
        import httpx

        try:
            async with httpx.AsyncClient(timeout=10.0, follow_redirects=True) as client:
                response = await client.get(uri, headers={"Range": "bytes=0-8191"})
            if response.status_code in (200, 206):
                status, detail = "reachable", f"Reachable · HTTP {response.status_code}"
                file_format = (config.get("file_format") or "").lower()
                if file_format == "csv":
                    first_line = response.text.splitlines()[0] if response.text else ""
                    header = [h.strip() for h in first_line.split(",") if h.strip()]
                    if header:
                        detected_fields = [{"field_name": h, "sample_value": None} for h in header]
                        detail += f" · {len(header)} columns detected in header"
            else:
                status, detail = "unreachable", f"Server responded HTTP {response.status_code}"
        except Exception as exc:
            status, detail = "unreachable", str(exc)[:500]
    else:
        status = "unverified"
        detail = (
            f"Reachability checks aren't supported yet for '{scheme or 'local path'}' references -- "
            "confirm manually that this path is reachable from the node before publishing."
        )

    with _cursor() as cur:
        cur.execute(
            """
            UPDATE dataset_source_config
            SET reachability_status = %s, reachability_checked_at = NOW(), reachability_detail = %s,
                detected_fields = COALESCE(%s, detected_fields), updated_at = NOW()
            WHERE dataset_id = %s
            RETURNING *
            """,
            (status, detail, Json(detected_fields) if detected_fields is not None else None, dataset_id),
        )
        return dict(cur.fetchone())


def save_database_reference_source(
    dataset_id: int, connection_string: str, table_name: str, engine: str
) -> dict:
    """Step 2 (database): record a read-only connection string + which
    table/view to register. Never connects here -- that's
    verify_database_reference, called explicitly via "Verify reachability"."""
    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        row = _upsert_source_config(
            cur, dataset_id,
            source_type="database",
            db_connection_string=connection_string,
            db_engine=engine,
            selected_table_name=table_name,
            reachability_status="unverified",
            reachability_checked_at=None,
            reachability_detail=None,
        )
    return _public_source_config(row)


def verify_database_reference(dataset_id: int) -> dict:
    """"Verify reachability" for a database reference -- PostgreSQL only
    this pass (see module docstring). Connects read-only, reads column
    names from information_schema, and never touches row data."""
    with _cursor() as cur:
        cur.execute("SELECT * FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,))
        config = cur.fetchone()
    if not config or not config.get("db_connection_string"):
        raise SourceConfigNotFoundError(dataset_id)

    engine = (config.get("db_engine") or "").strip().lower()
    detected_fields = None

    if engine not in ("postgresql", "postgres"):
        status = "unsupported"
        detail = f"Live verification for '{config.get('db_engine') or 'this engine'}' isn't built yet -- only PostgreSQL is supported this pass."
    else:
        import psycopg2

        schema, table = _split_schema_table(config.get("selected_table_name") or "")
        try:
            live_conn = psycopg2.connect(
                config["db_connection_string"],
                connect_timeout=5,
                options="-c default_transaction_read_only=on",
            )
            try:
                live_cur = live_conn.cursor(cursor_factory=RealDictCursor)
                live_cur.execute(
                    """
                    SELECT column_name, data_type FROM information_schema.columns
                    WHERE table_schema = %s AND table_name = %s
                    ORDER BY ordinal_position
                    """,
                    (schema, table),
                )
                columns = live_cur.fetchall()
            finally:
                live_conn.close()

            if not columns:
                status, detail = "unreachable", f"Connected, but table '{schema}.{table}' was not found or has no columns"
            else:
                detected_fields = [{"field_name": c["column_name"], "data_type": c["data_type"]} for c in columns]
                status, detail = "reachable", f"Connected read-only · {len(detected_fields)} columns detected"
        except Exception as exc:
            status, detail = "unreachable", str(exc)[:500]

    with _cursor() as cur:
        cur.execute(
            """
            UPDATE dataset_source_config
            SET reachability_status = %s, reachability_checked_at = NOW(), reachability_detail = %s,
                detected_fields = COALESCE(%s, detected_fields), updated_at = NOW()
            WHERE dataset_id = %s
            RETURNING *
            """,
            (status, detail, Json(detected_fields) if detected_fields is not None else None, dataset_id),
        )
        row = dict(cur.fetchone())
    return _public_source_config(row)


def save_map_service_source(
    dataset_id: int, map_service_url: str, layer_type: str, map_layer_name: str
) -> dict:
    """Step 2 (map service): record an already-hosted WMS/WFS/WMTS layer.
    This is a reference to the contributor's own GeoServer/MapServer --
    unrelated to map-module-backend, which hosts CPHR's own layers."""
    with _cursor() as cur:
        if not _dataset_row(cur, dataset_id):
            raise DatasetNotFoundError(dataset_id)
        return _upsert_source_config(
            cur, dataset_id,
            source_type="map_service",
            map_service_url=map_service_url,
            layer_type=layer_type,
            map_layer_name=map_layer_name,
            reachability_status="unverified",
            reachability_checked_at=None,
            reachability_detail=None,
        )


async def verify_map_service(dataset_id: int) -> dict:
    """"Verify reachability" for a map service reference: fetch
    GetCapabilities and confirm the named layer appears in it. The layer
    keeps rendering from the contributor's own server either way -- this
    only reads the capabilities/attribute-schema document."""
    with _cursor() as cur:
        cur.execute("SELECT * FROM dataset_source_config WHERE dataset_id = %s", (dataset_id,))
        config = cur.fetchone()
    if not config or not config.get("map_service_url"):
        raise SourceConfigNotFoundError(dataset_id)

    import httpx

    layer_name = config.get("map_layer_name")
    params = _capabilities_request_params(config.get("layer_type"))

    try:
        async with httpx.AsyncClient(timeout=10.0, follow_redirects=True) as client:
            response = await client.get(config["map_service_url"], params=params)
        if response.status_code == 200:
            body_lower = response.text.lower()
            if layer_name and layer_name.lower() not in body_lower:
                status = "reachable"
                detail = f"Reachable · GetCapabilities OK -- but layer '{layer_name}' wasn't found in it. Double-check the layer name."
            else:
                status = "reachable"
                detail = "Reachable · GetCapabilities OK" + (f" · layer '{layer_name}' found" if layer_name else "")
        else:
            status, detail = "unreachable", f"Server responded HTTP {response.status_code}"
    except Exception as exc:
        status, detail = "unreachable", str(exc)[:500]

    with _cursor() as cur:
        cur.execute(
            """
            UPDATE dataset_source_config
            SET reachability_status = %s, reachability_checked_at = NOW(), reachability_detail = %s, updated_at = NOW()
            WHERE dataset_id = %s
            RETURNING *
            """,
            (status, detail, dataset_id),
        )
        return dict(cur.fetchone())
